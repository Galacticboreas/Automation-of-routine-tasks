import os
from pprint import pprint
from app import ReportSettingsOrders
from openpyxl import load_workbook

config = ReportSettingsOrders()

cwd = os.getcwd()

pprint(f'мы находимся здесь: {cwd}')

os.chdir(config.path_dir)

wb = load_workbook('./Отчет заказы на производство от 13.02.2024 10ч00м.xlsx')

pprint(wb.get_sheet_names())
sheet = wb['Основной']
pprint(sheet.max_row)
pprint(sheet.max_column)
